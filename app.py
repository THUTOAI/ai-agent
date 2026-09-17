import io
import json
import os
import zipfile
from datetime import datetime
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from flask import Flask, jsonify, render_template, request

ROOT = Path(__file__).resolve().parent
DATA_DIR = ROOT / "data"
DATA_DIR.mkdir(exist_ok=True)

INVENTORY_FILE = DATA_DIR / "store_clothing_collection.xlsx"
SALES_FILE = DATA_DIR / "sales_records.json"
OPENROUTER_URL = "https://openrouter.ai/api/v1/chat/completions"
OPENROUTER_MODEL = os.environ.get(
    "OPENROUTER_MODEL", "nvidia/nemotron-3-nano-omni-30b-a3b-reasoning:free")

# Flask auto-discovers templates/ and static/ next to this file - no extra config needed.
app = Flask(__name__)


def read_inventory(workbook):
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError(
            "Install the Excel dependency with: pip install -r requirements.txt") from error

    try:
        sheet = load_workbook(workbook, read_only=True, data_only=True).active
    except zipfile.BadZipFile as error:
        raise ValueError(
            "The inventory file is not a valid .xlsx workbook. Replace it with an Excel file."
        ) from error
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []

    def key(value):
        return "".join(character for character in str(value or "").lower() if character.isalnum())

    headers = {key(value): index for index, value in enumerate(rows[0])}
    product_index = next((headers[name] for name in (
        "product", "productname", "name", "item") if name in headers), None)
    if product_index is None:
        raise ValueError(
            "Excel must include a Product, Product Name, Name, or Item column.")

    def value(row, names, default=""):
        index = next((headers[name]
                     for name in names if name in headers), None)
        return row[index] if index is not None and index < len(row) and row[index] is not None else default

    def list_value(row, names):
        raw = str(value(row, names, ""))
        return [item.strip() for item in raw.replace(";", ",").split(",") if item.strip()]

    inventory = []
    for row in rows[1:]:
        product = str(row[product_index] or "").strip()
        if not product:
            continue
        price = str(value(row, ("price", "pricezar", "cost"), "Contact us"))
        stock_quantity = value(
            row, ("stockqty", "quantity", "stockquantity"), None)
        stock = value(row, ("stock", "instock", "available"), True)
        if stock_quantity is None:
            stock_quantity = 1 if stock is True else 0
        try:
            stock_quantity = max(0, int(float(stock_quantity)))
        except (TypeError, ValueError):
            stock_quantity = 0
        if isinstance(stock, str):
            stock = stock.strip().lower() not in ("false", "no", "out", "out of stock", "0")
        image = str(value(row, ("image", "imagepath", "photo"), ""))
        inventory.append({
            "sku": str(value(row, ("sku", "productcode", "code"), "")),
            "category": str(value(row, ("category", "type"), "")),
            "name": product,
            "price": price if price.startswith("R") else f"R{price}",
            "stock": bool(stock) and stock_quantity > 0,
            "stock_quantity": stock_quantity,
            "image": image,
            "colors": list_value(row, ("colors", "colours", "color", "colour")),
            "sizes": list_value(row, ("sizes", "size")),
            "season": str(value(row, ("season",), "")),
            "material": str(value(row, ("material", "fabric"), "")),
        })
    return inventory


def inventory_payload():
    if not INVENTORY_FILE.exists():
        try:
            from inventory_data import INVENTORY
        except ImportError:
            return []
        return [{
            "sku": item.sku,
            "category": item.category,
            "name": item.product_name,
            "price": f"R{item.price_zar:.2f}",
            "stock": item.stock_quantity > 0,
            "stock_quantity": item.stock_quantity,
            "image": "",
            "colors": [item.color],
            "sizes": [item.size],
            "season": item.season,
            "material": item.material,
        } for item in INVENTORY]
    return read_inventory(INVENTORY_FILE)


def get_sales_records():
    if not SALES_FILE.exists():
        return []
    try:
        records = json.loads(SALES_FILE.read_text(encoding="utf-8"))
        return records if isinstance(records, list) else []
    except (OSError, json.JSONDecodeError):
        return []


def save_sales_records(records):
    SALES_FILE.write_text(json.dumps(records, indent=2), encoding="utf-8")


def stock_report():
    inventory = inventory_payload()
    sales = get_sales_records()
    sold_by_sku = {}
    for record in sales:
        sku = record.get("sku", "")
        sold_by_sku[sku] = sold_by_sku.get(
            sku, 0) + int(record.get("quantity", 0))
    report = []
    for item in inventory:
        sold = sold_by_sku.get(item["sku"], 0)
        current = max(0, item["stock_quantity"] - sold)
        target = max(10, sold * 2)
        report.append({
            **item,
            "stock": current > 0,
            "stock_quantity": current,
            "sold_quantity": sold,
            "reorder_quantity": max(0, target - current),
        })
    return report


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/inventory")
def api_inventory():
    try:
        return jsonify({"inventory": inventory_payload()})
    except (RuntimeError, ValueError) as error:
        return jsonify({"error": str(error)}), 500


@app.route("/api/stock-report")
def api_stock_report():
    try:
        return jsonify({"inventory": stock_report()})
    except (RuntimeError, ValueError) as error:
        return jsonify({"error": str(error)}), 500


@app.route("/api/chat", methods=["POST"])
def api_chat():
    try:
        request_data = request.get_json(force=True, silent=False)
        messages = request_data.get("messages", [])
        inventory = request_data.get("inventory", [])
        if not messages or not isinstance(messages, list):
            raise ValueError("At least one message is required")

        api_key = os.environ.get("OPENROUTER_API_KEY")
        if not api_key:
            return jsonify({"error": "OPENROUTER_API_KEY is not configured."}), 503

        system_message = {
            "role": "system",
            "content": (
                "You are Tumi, the friendly Streetcode Clothing shop assistant. "
                "Answer clearly and briefly. Only make product, price, stock, colour, "
                "and size claims using the supplied inventory. For anything outside "
                "the shop, say you can help with Streetcode Clothing questions. "
                "Never ask for or process payment details.\n\n"
                "Use only the customer catalogue supplied below when answering. "
                "Do not use the private owner stock report or invent products.\n\n"
                f"Customer catalogue:\n{json.dumps(inventory, ensure_ascii=True)}"
            ),
        }
        payload = json.dumps({
            "model": OPENROUTER_MODEL,
            "messages": [system_message, *messages[-10:]],
            "temperature": 0.4,
            "max_tokens": 250,
        }).encode("utf-8")
        upstream_request = Request(
            OPENROUTER_URL,
            data=payload,
            headers={
                "Authorization": f"Bearer {api_key}",
                "Content-Type": "application/json",
                "HTTP-Referer": "http://127.0.0.1:5000",
                "X-Title": "Streetcode Clothing chatbot",
            },
            method="POST",
        )
        with urlopen(upstream_request, timeout=30) as response:
            result = json.loads(response.read())
        answer = result["choices"][0]["message"]["content"].strip()
        return jsonify({"reply": answer})
    except (ValueError, KeyError, json.JSONDecodeError) as error:
        return jsonify({"error": str(error)}), 400
    except (HTTPError, URLError, TimeoutError) as error:
        return jsonify({"error": f"LLM request failed: {error}"}), 502


@app.route("/api/inventory/upload", methods=["POST"])
def api_inventory_upload():
    try:
        uploaded_file = request.files.get("file")
        if uploaded_file is None or uploaded_file.filename == "":
            raise ValueError("Choose an Excel file to upload.")
        content = uploaded_file.read()
        inventory = read_inventory(io.BytesIO(content))
        if not inventory:
            raise ValueError("The workbook has no product rows.")
        INVENTORY_FILE.write_bytes(content)
        return jsonify({"inventory": inventory, "message": f"Loaded {len(inventory)} products."})
    except (KeyError, RuntimeError, ValueError) as error:
        return jsonify({"error": str(error)}), 400


@app.route("/api/sales", methods=["POST"])
def api_sales():
    try:
        request_data = request.get_json(force=True, silent=False)
        sku = str(request_data.get("sku", "")).strip()
        quantity = int(request_data.get("quantity", 1))
        if not sku or quantity < 1:
            raise ValueError("A valid product and quantity are required.")
        item = next(
            (product for product in stock_report() if product["sku"] == sku), None)
        if item is None:
            raise ValueError("That product is not in the inventory.")
        if quantity > item["stock_quantity"]:
            raise ValueError(
                f"Only {item['stock_quantity']} unit(s) remain in stock.")
        records = get_sales_records()
        records.append({
            "sku": sku,
            "quantity": quantity,
            "recorded_at": datetime.now().isoformat(),
        })
        save_sales_records(records)
        return jsonify({"message": "Sale recorded.", "inventory": stock_report()})
    except (ValueError, TypeError, json.JSONDecodeError) as error:
        return jsonify({"error": str(error)}), 400


if __name__ == "__main__":
    port = int(os.environ.get("PORT", "5000"))
    app.run(host="0.0.0.0", port=port)