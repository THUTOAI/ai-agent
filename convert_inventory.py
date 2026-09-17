from __future__ import annotations

import argparse
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook


REQUIRED_COLUMNS = {
    "sku": "sku",
    "category": "category",
    "productname": "product_name",
    "color": "color",
    "size": "size",
    "season": "season",
    "material": "material",
    "pricezar": "price_zar",
    "stockqty": "stock_quantity",
}

STORE_INVENTORY_PRODUCTS = (
    ("STORE-TSHIRT", "T-Shirts", "Tshirt", 500.00,
     ("White",), ("S", "M", "L", "XL", "XXL")),
    ("STORE-PANTS", "Pants", "Pants", 450.00,
     ("Black", "Navy", "Grey", "Khaki"), ("S", "M", "L", "XL", "XXL")),
    ("STORE-JACKET", "Jackets", "Jacket", 650.00,
     ("Black",), ("S", "M", "L", "XL", "XXL")),
    ("STORE-BEANIES", "Headwear", "Beanies", 150.00,
     ("White", "Grey", "Blue", "Red", "Black", "Pink", "Yellow", "Green"), ()),
    ("STORE-SWEATER", "Sweaters", "Sweater", 650.00,
     ("Black",), ("S", "M", "L", "XL", "XXL")),
    ("STORE-CARDIGAN", "Sweaters", "Cardigan", 1200.00,
     ("Black", "White", "Red"), ("S", "M", "L", "XL", "XXL")),
    ("STORE-HOODIE", "Hoodies", "Hoodie", 750.00,
     ("Black", "Brown", "Red"), ("S", "M", "L", "XL", "XXL")),
    ("STORE-MAFIA-TRACKSUITS", "Tracksuits", "Mafia Tracksuits",
     1200.00, ("Black", "Navy"), ("S", "M", "L", "XL", "XXL")),
    ("STORE-SC-WOVEN-TRACKSUITS", "Tracksuits", "SC Woven Tracksuits",
     1000.00, ("Black", "White"), ("S", "M", "L", "XL", "XXL")),
    ("STORE-HANDBAG", "Accessories", "Handbag", 1000.00, ("Black", "Green"), ()),
    ("STORE-CAP", "Headwear", "Cap", 200.00,
     ("Black", "Brown", "White", "Red", "Blue"), ()),
)


def normalise_header(value: object) -> str:
    return "".join(character for character in str(value or "").lower() if character.isalnum())


def python_literal(value: object) -> str:
    if isinstance(value, bool):
        return repr(value)
    if isinstance(value, (int, float, Decimal)):
        return repr(value)
    return repr(str(value or "").strip())


def convert_workbook(source: Path, destination: Path) -> int:
    sheet = load_workbook(source, read_only=True, data_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError("The workbook is empty.")

    headers = {
        normalise_header(value): index for index, value in enumerate(rows[0])
    }
    missing = [name for name in REQUIRED_COLUMNS if name not in headers]
    if missing:
        raise ValueError(f"Missing columns: {', '.join(missing)}")

    products = []
    for row in rows[1:]:
        if not any(value is not None for value in row):
            continue
        values = {
            field: row[headers[column]] if headers[column] < len(row) else ""
            for column, field in REQUIRED_COLUMNS.items()
        }
        products.append(
            "    Product(\n"
            + "\n".join(
                f"        {field}={python_literal(value)},"
                for field, value in values.items()
            )
            + "\n    ),"
        )

    for sku, category, product_name, price, colors, sizes in STORE_INVENTORY_PRODUCTS:
        products.append(
            "    Product(\n"
            f"        sku={sku!r},\n"
            f"        category={category!r},\n"
            f"        product_name={product_name!r},\n"
            "        color='',\n"
            "        size='',\n"
            "        season='All Season',\n"
            "        material='Not specified',\n"
            f"        price_zar={price!r},\n"
            "        stock_quantity=0,\n"
            f"        available_colors={colors!r},\n"
            f"        available_sizes={sizes!r},\n"
            "        source='storeInventory',\n"
            "    ),"
        )

    output = '''from dataclasses import dataclass\n\n\n@dataclass(frozen=True)\nclass Product:\n    sku: str\n    category: str\n    product_name: str\n    color: str\n    size: str\n    season: str\n    material: str\n    price_zar: float\n    stock_quantity: int\n    available_colors: tuple[str, ...] = ()\n    available_sizes: tuple[str, ...] = ()\n    source: str = "workbook"\n\n\nINVENTORY = [\n'''
    output += "\n".join(products)
    output += "\n]\n"
    destination.write_text(output, encoding="utf-8")
    return len(products)


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Convert an inventory workbook to readable Python.")
    parser.add_argument("source", type=Path, nargs="?",
                        default=Path("data/store_clothing_collection.xlsx"))
    parser.add_argument("destination", type=Path, nargs="?",
                        default=Path("inventory_data.py"))
    args = parser.parse_args()
    count = convert_workbook(args.source, args.destination)
    print(f"Converted {count} products to {args.destination}")


if __name__ == "__main__":
    main()